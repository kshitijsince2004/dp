import sys
import os
import json
import argparse
import openpyxl

def generate_district_diary(template_path, output_path, scope_data, calc_data):
    wb = openpyxl.load_workbook(template_path)
    
    district_name = scope_data.get('self_name', 'DISTRICT').upper()
    cutoff_date = calc_data.get('cutoff_date', '')
    year_num = calc_data.get('year_num', 2026)
    children = scope_data.get('children', [])
    display_names = scope_data.get('display_names', {})
    
    # 1. Rcell DD (Keep sheet name intact so formulas referencing 'Rcell DD North' remain valid)
    for name in wb.sheetnames:
        if 'Rcell DD' in name:
            ws = wb[name]
            ws['B1'] = f"Rcell Crime Diary {year_num}"
            ws['A2'] = district_name
            for idx, ps in enumerate(children):
                r = 3 + idx
                ws.cell(row=r, column=1, value=display_names.get(ps['id'], ps['name']))
            break
            
    # 2. R Cell- Distt Crime
    if 'R Cell- Distt Crime' in wb.sheetnames:
        ws = wb['R Cell- Distt Crime']
        ws['A1'] = f"R Cell Daily Diary- {district_name} District"
        ws['B3'] = f"Upto Date {year_num - 1}"
        ws['E3'] = f"Upto Date {year_num}"
        
    # 3. E-FIR (Keep sheet name intact so formulas referencing 'E-FIR North,' remain valid)
    for name in wb.sheetnames:
        if 'E-FIR' in name:
            ws = wb[name]
            ws['A1'] = f"E-FIR OF {district_name} DISTRICT"
            break
            
    # 4. N-1,N-2,N-3
    if 'N-1,N-2,N-3 ' in wb.sheetnames:
        ws = wb['N-1,N-2,N-3 ']
        ws['A1'] = f"Daily {district_name} District Crime"
        
    # 5. D1,N-1,2,3 Res
    if ' D1,N-1,2,3 Res' in wb.sheetnames:
        ws = wb[' D1,N-1,2,3 Res']
        ws['A1'] = f"DAILY DIARY {district_name} DISTRICT"

    # 7. Daily Chart, Heinous, IPC
    if 'Daily Chart, Heinous, IPC' in wb.sheetnames:
        ws = wb['Daily Chart, Heinous, IPC']
        ws['A1'] = f"PS Wise Crime Chart — {district_name}"
        for idx, ps in enumerate(children):
            r = 6 + idx
            ws.cell(row=r, column=1, value=display_names.get(ps['id'], ps['name']))

    # 8. DCsP- Crime Chart
    if 'DCsP- Crime Chart ' in wb.sheetnames:
        ws = wb['DCsP- Crime Chart ']
        ws['A1'] = f"DAILY DIARY {district_name} DISTRICT"

    # 9. D-2 Heinous Brief Fact
    if 'D-2 Heinous Brief Fact' in wb.sheetnames:
        ws = wb['D-2 Heinous Brief Fact']
        ws['J1'] = cutoff_date
        heinous_list = calc_data.get('heinousList', [])
        for idx, item in enumerate(heinous_list):
            r = 4 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=item.get('ps_name', ''))
            ws.cell(row=r, column=3, value=item.get('fir_no', ''))
            ws.cell(row=r, column=9, value=item.get('brief_facts', ''))

    # 10. Upto PCR calls
    if 'Upto PCR calls 25-26' in wb.sheetnames:
        ws = wb['Upto PCR calls 25-26']
        for idx, ps in enumerate(children):
            r = 4 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=display_names.get(ps['id'], ps['name']))

    # 11. D-8 Brief Facts
    if 'D-8 Brief Facts' in wb.sheetnames:
        ws = wb['D-8 Brief Facts']
        ws['J1'] = cutoff_date
        ws['B2'] = f"DAILY CRIME (PS FIR ) WITH BRIEF FACTS OF {district_name} DISTRICT  :"
        full_fir = calc_data.get('fullFirList', [])
        for idx, item in enumerate(full_fir):
            r = 5 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=item.get('ps_name', ''))
            ws.cell(row=r, column=3, value=item.get('fir_no', ''))
            ws.cell(row=r, column=8, value=item.get('brief_facts', ''))

    # 12. D-9 FIR Arrests
    if 'D-9 FIR Arrests ' in wb.sheetnames:
        ws = wb['D-9 FIR Arrests ']
        ws['C2'] = f"DAILY MORNING DIARY REGARDING PERSONS ARRESTED IN FIR, {district_name} DISTRICT"
        fir_arrests = calc_data.get('firArrestsList', [])
        for idx, item in enumerate(fir_arrests):
            r = 6 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=item.get('ps_name', ''))
            ws.cell(row=r, column=3, value=item.get('person_name', ''))
            ws.cell(row=r, column=4, value=item.get('age', ''))
            ws.cell(row=r, column=7, value=item.get('fir_no', ''))
            ws.cell(row=r, column=13, value=item.get('custody_status', ''))

    # 13. D-9 Kal Arrests
    if 'D-9 Kal Arrests' in wb.sheetnames:
        ws = wb['D-9 Kal Arrests']
        ws['A2'] = f"DAILY MORNING DIARY REGARDING PERSONS ARRESTED IN KALANDRAS,  {district_name} DISTRICT"
        kal_arrests = calc_data.get('kalArrestsList', [])
        for idx, item in enumerate(kal_arrests):
            r = 6 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=item.get('ps_name', ''))
            ws.cell(row=r, column=3, value=item.get('person_name', ''))
            ws.cell(row=r, column=4, value=item.get('age', ''))
            ws.cell(row=r, column=13, value=item.get('custody_status', ''))

    # 14. D10 Action of 66 DP Act
    if 'D10 Action of 66 DP Act ' in wb.sheetnames:
        ws = wb['D10 Action of 66 DP Act ']
        for idx, ps in enumerate(children):
            r = 4 + idx
            ws.cell(row=r, column=1, value=display_names.get(ps['id'], ps['name']))
            ws.cell(row=r, column=11, value=display_names.get(ps['id'], ps['name']))

    # 15. D13 66DP
    if 'D13 66DP' in wb.sheetnames:
        ws = wb['D13 66DP']
        ws['A2'] = f"VEHICLE SEIZED OF {district_name} DISTRICT UNDER 66 DP ACT"
        for idx, ps in enumerate(children):
            r = 5 + idx
            ws.cell(row=r, column=1, value=display_names.get(ps['id'], ps['name']))

    # 16. G-22 Daily Crime
    if 'G-22 Daily Crime' in wb.sheetnames:
        ws = wb['G-22 Daily Crime']
        ws['A1'] = f"Daily Crime {district_name} District"
        ws['A5'] = district_name

    # 17. Accident Cases
    if 'Accident Cases' in wb.sheetnames:
        ws = wb['Accident Cases']
        ws['A3'] = district_name
        accidents = calc_data.get('accidentList', [])
        for idx, item in enumerate(accidents):
            r = 6 + idx
            ws.cell(row=r, column=1, value=idx + 1)
            ws.cell(row=r, column=2, value=f"{item.get('ps_name', '')} - {item.get('fir_no', '')}: {item.get('brief_facts', '')}")

    # Ensure output directory exists before saving
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    wb.save(output_path)
    print("SUCCESS")

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--payload', required=True)
    args = parser.parse_args()

    with open(args.payload, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    generate_district_diary(
        args.template,
        args.output,
        payload.get('scope', {}),
        payload.get('calcData', {})
    )
