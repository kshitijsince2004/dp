#!/usr/bin/env python3
"""Stage 1 — verify ref.* natural keys (DB_SCHEMA.md §8 ⚠ items) against Menu_Tables.xlsx.

Read-only. Prints a verdict per ⚠ key and drafts the heinous overlay.
"""
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]  # backend/
XLSX = ROOT / "seeds" / "Menu_Tables.xlsx"
EXPORTS = ROOT.parent / "docs" / "db-audit" / "exports"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v in ("\\N", "NULL", ""):
            return None
    return v


def clean_int(v):
    v = clean(v)
    if v is None:
        return None
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return None


def rows_of(sheet_name, skip_header=1):
    ws = wb[sheet_name]
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= skip_header:
            continue
        out.append(r)
    return out


def report_unique(label, values, allow_null=False):
    """values: list of (row_number, key). Returns set of non-null keys."""
    nulls = [rn for rn, k in values if k is None]
    nonnull = [(rn, k) for rn, k in values if k is not None]
    dupes = {k: n for k, n in Counter(k for _, k in nonnull).items() if n > 1}
    verdict = "UNIQUE" if not dupes else f"NOT UNIQUE ({len(dupes)} duplicated keys)"
    nullnote = f", {len(nulls)} NULL" if nulls else ""
    print(f"  {label}: {len(nonnull)} non-null rows{nullnote} -> {verdict}")
    if dupes:
        sample = list(dupes.items())[:8]
        print(f"    dup sample: {sample}")
    if nulls and not allow_null:
        print(f"    NULL rows at: {nulls[:10]}{'...' if len(nulls) > 10 else ''}")
    return {k for _, k in nonnull}


print(f"Workbook: {XLSX}\n")

# ── acts ─────────────────────────────────────────────────────────────────────
print("[ref.acts]")
acts = [(i + 2, clean_int(r[0])) for i, r in enumerate(rows_of("act"))]
acts = [(rn, k) for rn, k in acts if k is not None]
act_cds = report_unique("act_cd", acts)

# ── sections ─────────────────────────────────────────────────────────────────
print("[ref.sections]  (⚠ act_sec_cd; also section_code as FK target of mapping)")
sec_rows = []
for i, r in enumerate(rows_of("section")):
    section_code = clean(r[0])
    if section_code is None:
        continue
    sec_rows.append((i + 2, section_code, clean(r[2]), clean(r[3])))  # code, section_cd, act_sec_cd
section_codes = report_unique("section_code", [(rn, sc) for rn, sc, _, _ in sec_rows])
report_unique("section_cd", [(rn, scd) for rn, _, scd, _ in sec_rows], allow_null=True)
report_unique("act_sec_cd", [(rn, ascd) for rn, _, _, ascd in sec_rows], allow_null=True)

# ── major heads ──────────────────────────────────────────────────────────────
print("[ref.major_heads]")
majors = [(i + 2, clean_int(r[0])) for i, r in enumerate(rows_of("major head"))]
majors = [(rn, k) for rn, k in majors if k is not None]
major_cds = report_unique("major_head_code", majors)

# ── minor heads ──────────────────────────────────────────────────────────────
print("[ref.minor_heads]  (⚠ composite (major_head_code, minor_head_cd))")
minors = []
for i, r in enumerate(rows_of("minor head")):
    cd = clean_int(r[0])
    if cd is None:
        continue
    minors.append((i + 2, cd, clean_int(r[2])))
report_unique("minor_head_cd alone", [(rn, cd) for rn, cd, _ in minors])
report_unique("(major_head_code, minor_head_cd)", [(rn, (mj, cd)) for rn, cd, mj in minors])
bad_fk = [(rn, mj) for rn, _, mj in minors if mj not in major_cds]
print(f"  FK major_head_code -> major_heads: {len(bad_fk)} unresolved" + (f" e.g. {bad_fk[:5]}" if bad_fk else ""))

# ── major_minor_mapping ──────────────────────────────────────────────────────
print("[ref.major_minor_mapping]  (⚠ sec_mjrhd_cd; FKs act_cd/section_code/major_head_code)")
mmm = []
for i, r in enumerate(rows_of("Major_Minor_Mapping")):
    cd = clean_int(r[0])
    if cd is None:
        continue
    mmm.append((i + 2, cd, clean_int(r[2]), clean(r[3]), clean_int(r[4])))
report_unique("sec_mjrhd_cd", [(rn, cd) for rn, cd, *_ in mmm])
report_unique("(act_cd, section_code) pair", [(rn, (a, s)) for rn, _, a, s, _ in mmm])
for label, idx, universe in (("act_cd->acts", 2, act_cds), ("section_code->sections", 3, section_codes), ("major_head_code->majors", 4, major_cds)):
    bad = [(row[0], row[idx]) for row in mmm if row[idx] is not None and row[idx] not in universe]
    nulls = sum(1 for row in mmm if row[idx] is None)
    print(f"  FK {label}: {len(bad)} unresolved, {nulls} NULL" + (f" e.g. {bad[:5]}" if bad else ""))

# ── local heads ──────────────────────────────────────────────────────────────
print("[ref.local_heads]")
locals_ = []
for i, r in enumerate(rows_of("local head")):
    cd = clean_int(r[0])
    if cd is None:
        continue
    locals_.append((i + 2, cd, clean(r[2])))
report_unique("local_head_cd", [(rn, cd) for rn, cd, _ in locals_])

# ── property types / OTHERS ──────────────────────────────────────────────────
print("[ref.property_types]  (⚠ parent_cd)")
ptypes = []
for i, r in enumerate(rows_of("Property Type")):
    srno = clean_int(r[0])
    if srno is None:
        continue
    ptypes.append((i + 2, srno, clean_int(r[2])))
report_unique("parent_srno", [(rn, s) for rn, s, _ in ptypes])
ptype_cds = report_unique("parent_cd", [(rn, c) for rn, _, c in ptypes])

print("[ref.other_property_categories + items]  (⚠ parent_cd, property_cd — OTHERS sheet, 2 sections)")
cat_rows, item_rows = [], []
state = None
for i, r in enumerate(wb["OTHERS"].iter_rows(values_only=True), start=1):
    c0 = clean(r[0])
    if c0 == "parent_srno":
        state = "MAJOR"; continue
    if c0 == "property_cd":
        state = "MINOR"; continue
    if state == "MAJOR":
        srno, cd = clean_int(r[0]), clean_int(r[2])
        if srno is not None and cd is not None:
            cat_rows.append((i, srno, cd))
    elif state == "MINOR":
        pcd, parent = clean_int(r[0]), clean_int(r[2])
        if pcd is not None and parent is not None:
            item_rows.append((i, pcd, parent, clean(r[4])))
cat_cds = report_unique("categories.parent_cd", [(rn, cd) for rn, _, cd in cat_rows])
report_unique("items.property_cd", [(rn, cd) for rn, cd, _, _ in item_rows])
bad = [(rn, p) for rn, _, p, _ in item_rows if p not in cat_cds]
print(f"  FK items.parent_cd -> categories: {len(bad)} unresolved" + (f" e.g. {bad[:5]}" if bad else ""))
overlap = ptype_cds & cat_cds
print(f"  property_types.parent_cd ∩ other_categories.parent_cd = {sorted(overlap) if overlap else 'none'} (tables are separate; overlap OK but noted)")

# ── ARMS sheet: 4 sections ───────────────────────────────────────────────────
print("[ARMS AND AMMUNITION — 4-section parse]  (⚠ fire_arms_subtypes)")
made, cats, fire, subs = [], [], [], []
state = None
for i, r in enumerate(wb["ARMS AND AMMUNITION"].iter_rows(values_only=True), start=1):
    c0 = clean(r[0])
    if c0 == "arms_made_cd":
        state = "MADE"; continue
    if c0 == "arms_category_cd":
        state = "CAT"; continue
    if c0 == "fire_arms_cd":
        state = "FIRE"; continue
    if c0 == "arms_subtype_cd":
        state = "SUB"; continue
    if isinstance(c0, str) and c0 and clean_int(c0) is None:
        # section title rows like 'Made', 'Category', ...
        continue
    if state == "MADE":
        cd = clean_int(r[0])
        if cd is not None:
            made.append((i, cd, clean(r[2])))
    elif state == "CAT":
        cd = clean_int(r[0])
        if cd is not None:
            cats.append((i, cd, clean(r[2])))
    elif state == "FIRE":
        cd = clean_int(r[0])
        if cd is not None:
            fire.append((i, cd, clean_int(r[2]), clean(r[3])))
    elif state == "SUB":
        cd = clean_int(r[0])
        if cd is not None:
            subs.append((i, cd, clean_int(r[2]), clean(r[3])))
print(f"  sections found: made={len(made)}, categories={len(cats)}, fire_arms={len(fire)}, subtypes={len(subs)}")
report_unique("arms_made_cd", [(rn, cd) for rn, cd, _ in made])
cat_ids = report_unique("arms_category_cd", [(rn, cd) for rn, cd, _ in cats])
fire_ids = report_unique("fire_arms_cd", [(rn, cd) for rn, cd, _, _ in fire])
report_unique("arms_subtype_cd", [(rn, cd) for rn, cd, _, _ in subs])
bad = [(rn, c) for rn, _, c, _ in fire if c not in cat_ids]
print(f"  FK fire_arms.arms_category_cd -> categories: {len(bad)} unresolved" + (f" e.g. {bad[:5]}" if bad else ""))
bad = [(rn, c) for rn, _, c, _ in subs if c not in fire_ids]
print(f"  FK subtypes.arms_type_cd -> fire_arms: {len(bad)} unresolved" + (f" e.g. {bad[:8]}" if bad else ""))

# ── simple lookups ───────────────────────────────────────────────────────────
for sheet, key in [("AUTOMOBILES AND OTHERS", "automobile_cd"), ("COIN AND CURRENCY", "currency_type_cd"),
                   ("CULTURAL PROPERTY", "cultural_prop_cd"), ("Documents", "document_type_cd"),
                   ("DRUGS NARCOTIC", "drug_type_cd"), ("ELECTRICAL AND ELECTRONIC GOODS", "electric_goods_cd"),
                   ("EXPLOSIVES", "explosive_type_cd"), ("JEWELLERY", "jewelry_type_cd")]:
    vals = [(i + 2, clean_int(r[0])) for i, r in enumerate(rows_of(sheet))]
    vals = [(rn, k) for rn, k in vals if k is not None]
    print(f"[{sheet}]")
    report_unique(key, vals)

# ── beats ────────────────────────────────────────────────────────────────────
print("[ref.beats]  ((ps_cd, beat_cd) + ps_cd resolution)")
beats = []
for i, r in enumerate(rows_of("Beat")):
    bcd = clean(r[0])
    if bcd is None:
        continue
    beats.append((i + 2, str(bcd), clean(r[3]) and str(clean(r[3]))))
report_unique("beat_cd alone", [(rn, b) for rn, b, _ in beats])
report_unique("(ps_cd, beat_cd)", [(rn, (p, b)) for rn, b, p in beats])
ps_cds = Counter(p for _, _, p in beats)
print(f"  distinct ps_cd in Beat sheet: {len(ps_cds)}")
hier = json.load(open(EXPORTS / "hierarchy_nodes.json"))
ps_codes_db = {n["code"] for n in hier if n["node_type"] == "PS"}
resolved = {p for p in ps_cds if p in ps_codes_db}
print(f"  ps_cd resolvable against live hierarchy codes: {len(resolved)}/{len(ps_cds)}"
      f"  (live PS codes look like: {sorted(ps_codes_db)[:2]} — official numeric codes NOT in DB/repo)")

# ── heinous overlay draft ────────────────────────────────────────────────────
print("[heinous overlay draft]")
HEINOUS = ["Murder", "Attempt to murder", "Rape", "Gang rape", "Kidnapping for ransom",
           "Robbery", "Dacoity", "Acid attack", "Terrorism-related offences"]  # exact list from old seeds/02_menu_tables.js


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower()) if s else ""


matched, unmatched = {}, []
for h in HEINOUS:
    hits = [(cd, name) for _, cd, name in locals_ if norm(h) in norm(name) or norm(name) in norm(h)]
    if hits:
        matched[h] = hits
    else:
        unmatched.append(h)
print(f"  matched {len(matched)}/{len(HEINOUS)} heinous names against {len(locals_)} local heads")
for h, hits in matched.items():
    print(f"    {h}: {[(cd, nm) for cd, nm in hits[:4]]}{' …' if len(hits) > 4 else ''}")
print(f"  UNMATCHED: {unmatched}")
