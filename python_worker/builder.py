import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_HEADER_FILL  = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
_HEADER_FONT  = Font(name='Arial', size=9, bold=True, color='FFFFFF')
_HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_DATA_FONT    = Font(name='Arial', size=9)
_DATA_ALIGN   = Alignment(horizontal='left', vertical='center', wrap_text=True)
_ZEBRA_FILL   = PatternFill(start_color='EBF3FA', end_color='EBF3FA', fill_type='solid')
_THIN_BORDER  = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)
_TITLE_FONT   = Font(name='Arial', size=11, bold=True, color='1F3864')
_META_FONT    = Font(name='Arial', size=9, italic=True, color='595959')


TABLE_NAME_TO_TEMPLATE_TITLE = {
    'excel_1manual_fir': '1. Manual FIR',
    'excel_2eburglary_cases': '2. E-Burglary Cases',
    'excel_3ehouse_theft_cases': '3. E-House Theft Cases',
    'excel_4eother_theft_cases': '4. E-Other Theft Cases',
    'excel_5mvt_cases': '5. MVT Cases',
    'excel_8arrested_kalandara': '6. Arrested - Kalandara  Preven',
    'excel_9arrested_efir_theft': '7. Arrested - E-FIR Theft',
    'excel_7arrested_east_district': 'Arrested - District',
    'excel_10arrested_efir_mv_theft': '8. Arrested - E-FIR MV Theft',
    'excel_13arrested_24_hrs_list': '11. Arrested - Last 24 Hrs',
    'excel_14pi_disposal_manual': '12. PI Disposal - Manual',
    'excel_15pi_disposal_eproperty': '13. PI Disposal - E-Theft',
    'excel_16pi_disposal_emvt': '14. PI Disposal - E-MVT',
    'excel_18missing_persons': '15. Missing Persons',
    'excel_19uidb': '16. UIDB (Unidentified Bodies)',
    'excel_20abandoned_persons': '17. Abandoned Persons',
    'excel_21traced_persons': '18. Traced Persons',
    'excel_25inquest_registered': '19. Inquest Registered',
    'excel_26inquest_acpsdm_disposal': '20. Inquest ACPSDM Disposal',
    'excel_28fir_goswara_summary': '21. FIR Goswara Summary',
}


def _resolve_header(col_key, sheet_column_labels, shared_labels):
    """Look up display label: sheet overrides > shared labels > title-case fallback."""
    if col_key in sheet_column_labels:
        return sheet_column_labels[col_key]
    if col_key in shared_labels:
        return shared_labels[col_key]
    return col_key.replace('_', ' ').title()


SHARED_COLUMN_LABELS = {
    'sn': 'S.N.', 'sno': 'S.No.', 's_no': 'S. No.', 'sr_no': 'Sr. No.', 'sr': 'Sr.',
    'ps': 'Police Station',
    'fir_no': 'FIR No.', 'efir_no': 'E-FIR No.',
    'dd_nofir_no': 'DD No./FIR No.', 'firdd_no': 'FIR / DD No.',
    'dd_no': 'DD No.', 'dd_date': 'DD Date',
    'us': 'U/S (Act + Section)',
    'io': 'Name of IO (Rank / Name / PIS No.)', 'name_of_io': 'Name of IO (Rank / Name / PIS No.)', 'io_name': 'IO Name',
    'io_mobile_no': 'IO Mobile No.',
    'rank_of_io': 'Rank of IO', 'mobile_no_of_io': 'Mobile No. of IO',
    'complainant_details': 'Complainant (Name / S/O / R/O Address)',
    'arrested_details': 'Arrested Person (Name / Age / S/O / R/O Address)',
    'accused_details': 'Accused (Name / Age / S/O / R/O Address)',
    'po_details': 'PO Details (Name / S/O / R/O Address)',
    'deceased_details': 'Deceased (Name / Age / S/O / R/O Address)',
    'traced_person_details': 'Traced Person (Name / S/O / R/O Address)',
    'place_of_occurrence': 'Place of Occurrence (House No. / Street / Colony / Village / City / Tehsil / Landmark / District)',
    'place_of_occurrence_1': 'Place of Occurrence (2)',
    'time_of_occurrence': 'Date & Time of Occurrence (DD/MM/YYYY HH:MM)',
    'date_of_occurrence': 'Date & Time of Occurrence (DD/MM/YYYY HH:MM)',
    'pcjcbail': 'Custody Status (PC / JC / Bail)',
    'recovery': 'Recovery(Details of property recovered)',
    'beat_no': 'Beat No.',
    'stolen_items': 'Stolen Property (Details of the stolen property in the case)',
    'cause_of_death': 'Cause of Death',
    'challan_untrace_cancel': 'Disposal (Challan / Untrace / Cancel)',
    'name_of_operator_to_whom_mps': 'Name of Operator (MPS)',
    'found_place': 'Found Place (House No. / Street / Colony / Village / City / Tehsil / Landmark / District)',
    'found_date': 'Found Date',
    'upper_dress_color': 'Upper Dress Color',
    'lower_dress_color': 'Lower Dress Color',
    'prev_involvement_no_of_cases': 'Prev. Involvement (Y/N)',
    'whether_accused_is_bc_or_not': 'Accused BC(Y/N)',
    'group_patrolling': 'Group Patrolling',
    'cycle_patrolling': 'Cycle Patrolling',
    'by_antisnatching_team': 'By Anti-Snatching Team',
    'by_prahari': 'By Prahari',
    'by_eyes_ears_scheme_members': 'By Eyes & Ears Scheme Members',
    'pcr_call': 'PCR Call',
}


def build_workbook(sheets_data, sheets, file_path, date=''):
    """
    Build/populate daily-diary Excel workbook based directly on reference template Daily_Diary_16Jul2026_AllStations.xlsx.

    sheets_data: { table_name: [row_dict, ...] }
    sheets:      list of sheet descriptor dicts from registry.SHEETS
    file_path:   absolute path for the output .xlsx file
    date:        optional date string shown in sheet headers
    """
    template_file = os.path.join(os.path.dirname(__file__), 'templates', 'Daily_Diary_16Jul2026_AllStations.xlsx')
    if not os.path.exists(template_file):
        template_file = os.path.abspath('./Daily_Diary_16Jul2026_AllStations (1).xlsx')

    if os.path.exists(template_file):
        wb = openpyxl.load_workbook(template_file)
        kept_sheet_titles = set()

        for sheet_def in sheets:
            table_name = sheet_def['table_name']
            col_keys = sheet_def['columns']
            rows = sheets_data.get(table_name, [])

            target_title = TABLE_NAME_TO_TEMPLATE_TITLE.get(table_name)
            if not target_title or target_title not in wb.sheetnames:
                sheet_label = sheet_def.get('label') or table_name.replace('_', ' ').title()
                ws = wb.create_sheet(title=sheet_label[:31])
                kept_sheet_titles.add(ws.title)

                ws.cell(row=1, column=1, value=sheet_label).font = _TITLE_FONT
                ws.cell(row=2, column=1, value=f"Date: {date}" if date else "").font = _META_FONT
                ws.cell(row=3, column=1, value="")

                col_labels = [SHARED_COLUMN_LABELS.get(k, k.replace('_', ' ').title()) for k in col_keys]
                ws.append(col_labels)
                header_row_idx = 4
                for ci in range(1, len(col_labels) + 1):
                    c = ws.cell(row=header_row_idx, column=ci)
                    c.fill = _HEADER_FILL
                    c.font = _HEADER_FONT
                    c.alignment = _HEADER_ALIGN
                    c.border = _THIN_BORDER

                if rows:
                    for row_offset, row_dict in enumerate(rows, start=1):
                        row_values = [row_dict.get(k, '') for k in col_keys]
                        ws.append(row_values)
                        actual_row = header_row_idx + row_offset
                        ws.row_dimensions[actual_row].height = 18
                        use_zebra = (row_offset % 2 == 0)
                        for ci, val in enumerate(row_values, start=1):
                            cell = ws.cell(row=actual_row, column=ci)
                            cell.font = _DATA_FONT
                            cell.alignment = _DATA_ALIGN
                            cell.border = _THIN_BORDER
                            if use_zebra:
                                cell.fill = _ZEBRA_FILL
                continue

            kept_sheet_titles.add(target_title)
            ws = wb[target_title]

            # Update date in metadata row if present
            if date and ws.cell(row=2, column=1).value:
                cur_val = str(ws.cell(row=2, column=1).value)
                if 'Date:' in cur_val or 'Period:' in cur_val:
                    ws.cell(row=2, column=1).value = f"Date: {date}"

            # Clear template sample data rows (row 5 downwards)
            if ws.max_row > 4:
                ws.delete_rows(5, ws.max_row - 4)

            header_row_idx = 4

            if rows:
                for row_offset, row_dict in enumerate(rows, start=1):
                    row_values = [row_dict.get(k, '') for k in col_keys]
                    ws.append(row_values)
                    actual_row = header_row_idx + row_offset
                    ws.row_dimensions[actual_row].height = 18
                    use_zebra = (row_offset % 2 == 0)
                    for ci, val in enumerate(row_values, start=1):
                        cell = ws.cell(row=actual_row, column=ci)
                        cell.font = _DATA_FONT
                        cell.alignment = _DATA_ALIGN
                        cell.border = _THIN_BORDER
                        if use_zebra:
                            cell.fill = _ZEBRA_FILL

        # Remove any template sheets that were not in active_tables/sheets
        for sheet_title in list(wb.sheetnames):
            if sheet_title not in kept_sheet_titles and len(wb.worksheets) > 1:
                wb.remove(wb[sheet_title])

        wb.save(file_path)
        print(f'[Builder] Template-populated workbook saved: {file_path} ({len(wb.worksheets)} sheets)')
        return

    # Fallback to scratch builder if template file is unavailable
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_def in sheets:
        table_name   = sheet_def['table_name']
        label        = sheet_def['label']
        num          = sheet_def['num']
        col_keys     = sheet_def['columns']
        col_labels   = sheet_def.get('column_labels', {})

        rows = sheets_data.get(table_name, [])

        if label == 'Arrested - District' or table_name == 'excel_7arrested_east_district':
            sheet_title = 'Arrested - District'
        else:
            sheet_title = f'{num}. {label}'
        for bad_char in ':\\/?*[]':
            sheet_title = sheet_title.replace(bad_char, '')
        sheet_title = sheet_title[:31]

        ws = wb.create_sheet(title=sheet_title)

        ws.append([f'PHAROS Daily Diary — {label}'])
        ws.cell(row=ws.max_row, column=1).font = _TITLE_FONT
        ws.append([f'Date: {date}'] if date else [''])
        ws.cell(row=ws.max_row, column=1).font = _META_FONT
        ws.append([])

        col_headers = [_resolve_header(k, col_labels, SHARED_COLUMN_LABELS) for k in col_keys]
        ws.append(col_headers if col_headers else ['No Data'])
        header_row_idx = ws.max_row
        ws.row_dimensions[header_row_idx].height = 22
        for ci in range(1, len(col_headers) + 1):
            cell = ws.cell(row=header_row_idx, column=ci)
            cell.fill   = _HEADER_FILL
            cell.font   = _HEADER_FONT
            cell.alignment = _HEADER_ALIGN
            cell.border = _THIN_BORDER

        if not rows:
            ws.append(['No records for this date.'])
        else:
            for row_offset, row_dict in enumerate(rows, start=1):
                row_values = [row_dict.get(k, '') for k in col_keys]
                ws.append(row_values)
                actual_row = header_row_idx + row_offset
                ws.row_dimensions[actual_row].height = 18
                use_zebra = (row_offset % 2 == 0)
                for ci, val in enumerate(row_values, start=1):
                    cell = ws.cell(row=actual_row, column=ci)
                    cell.font      = _DATA_FONT
                    cell.alignment = _DATA_ALIGN
                    cell.border    = _THIN_BORDER
                    if use_zebra:
                        cell.fill = _ZEBRA_FILL

        for col_cells in ws.columns:
            max_len = 10
            for cell in col_cells:
                if cell.row < header_row_idx:
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = openpyxl.utils.get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

    if not wb.worksheets:
        ws = wb.create_sheet(title='No Data')
        ws.append(['No records found for the selected date and filters.'])

    wb.save(file_path)
    print(f'[Builder] Fallback workbook saved: {file_path} ({len(wb.worksheets)} sheets)')

